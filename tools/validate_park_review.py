"""Validate a newly generated blank PK-02E review workbook; never changes it.
Not an importer for completed human decisions or production coordinates.
"""
import argparse,collections,csv,hashlib,json,re,zipfile
from pathlib import Path
from urllib.parse import urlsplit,parse_qs
import openpyxl
from park_review_workbook import HEADERS,PRIORITY,STATUSES,DECISIONS,PROTECTED_DECISIONS,TYPES,valid,sha,load_source,prepare

def validate(path,receipt,source=None,supplement=None):
 path=Path(path);before=sha(path);w=openpyxl.load_workbook(path,data_only=False,keep_links=False);checks=[]
 def check(name,condition):
  if not condition:raise AssertionError(name)
  checks.append({'check':name,'passed':True})
 try:
  main=w['審核表'];risk=w['高風險優先審核'];stats=w['審核統計'];history=w['既有人工修正紀錄']
  check('exact_headers',list(next(main.values))==HEADERS)
  rows=list(main.iter_rows(min_row=2));ids=[r[0].value for r in rows];byid={r[0].value:r for r in rows}
  check('635_rows_and_unique_exact_np_ids',len(rows)==635 and len(byid)==635 and set(ids)=={f'NP-{i}' for i in range(1,636)})
  check('priority_first_eight_exact_order',ids[:8]==PRIORITY)
  check('all_pending_no_decisions_or_corrections',all(r[14].value=='待審' and all(r[i].value in (None,'') for i in [15,16,17,19,20,21,22]) and r[18].value=='未知' for r in rows))
  protected={r[0].value for r in rows if r[23].value=='保護'}
  check('nine_protected_exact_receipt_ids',len(protected)==9 and protected==set(receipt['protectedIds']))
  check('all_source_cells_locked',all(c.protection.locked for r in rows for c in list(r[:14])+list(r[23:])))
  check('nine_correction_pairs_locked',all(byid[id][16].protection.locked and byid[id][17].protection.locked for id in protected))
  check('unprotected_input_fields_editable',all(not r[i].protection.locked for r in rows for i in range(14,23) if not (r[0].value in protected and i in [16,17])))
  valid_links=0;invalid=[]
  for r in rows:
   id=r[0].value;cell=r[12];candidate=valid(r[6].value,r[7].value)
   if candidate:
    check(f'{id}_candidate_clickable',cell.hyperlink is not None and cell.hyperlink.target==cell.value)
    u=urlsplit(cell.value);q=parse_qs(u.query);check(f'{id}_candidate_url',u.scheme=='https' and u.netloc=='www.google.com' and u.path=='/maps/search/' and q.get('api')==['1'])
    lat,lon=map(float,q['query'][0].split(','));check(f'{id}_candidate_query_matches_id_values',abs(lat-float(r[6].value))<1e-12 and abs(lon-float(r[7].value))<1e-12);valid_links+=1
   else:check(f'{id}_invalid_no_coordinate_link',cell.value is None and cell.hyperlink is None);invalid.append(id)
   cell=r[13];u=urlsplit(cell.value);q=parse_qs(u.query)
   check(f'{id}_encoded_name_url',u.scheme=='https' and u.netloc=='www.google.com' and u.path=='/maps/search/' and q.get('api')==['1'] and q.get('query')==[f'{r[1].value} 新北市 {r[2].value}'] and cell.hyperlink.target==cell.value and cell.value.isascii())
  check('634_valid_links_and_np480_invalid',valid_links==634 and invalid==['NP-480'])
  cids={r[0].value for r in rows if r[9].value=='C'};rids=[r[0].value for r in risk.iter_rows(min_row=2)]
  check('62_risk_ids_unique_same_master',len(rids)==62 and len(set(rids))==62 and set(rids)==cids)
  memberships={id:set(filter(None,(r[24].value or '').split('、'))) for id,r in byid.items()};groupids=set().union(*memberships.values());dups={id for id,g in memberships.items() if g}
  check('23_groups_54_unique_no_double_count',len(groupids)==23 and len(dups)==54)
  check('eight_and_duplicates_union_62_not_double_count',set(PRIORITY)|dups==cids)
  for row in risk.iter_rows(min_row=2):
   master=byid[row[0].value][0].row
   check(f'{row[0].value}_risk_internal_link',row[0].hyperlink.location==f"'審核表'!A{master}" and row[0].hyperlink.target is None)
   if valid(row[6].value,row[7].value):check(f'{row[0].value}_risk_coordinate_clickable',row[12].hyperlink is not None)
   for i in range(14,23):
    col=row[i].column_letter;expected=f'=IF(\'審核表\'!{col}{master}="","",\'審核表\'!{col}{master})'
    check(f'{row[0].value}_risk_formula_{col}',row[i].value==expected and row[i].protection.locked)
  check('filter_and_freeze',main.auto_filter.ref=='A1:AB636' and main.freeze_panes=='C2' and main.protection.sheet and main.protection.autoFilter is False)
  check('conditional_formats_present',len(main.conditional_formatting)>=4)
  rules=list(main.data_validations.dataValidation)
  check('eight_validations_stop_errors',len(rules)==8 and all(r.showErrorMessage and r.errorStyle=='stop' for r in rules))
  check('named_dropdowns_resolve',all(n in w.defined_names for n in ['ReviewStatuses','ReviewDecisions','PositionTypes','ProtectedDecisions']))
  for col,expected in [('A',STATUSES),('B',DECISIONS),('C',TYPES),('D',PROTECTED_DECISIONS)]:check('dropdown_values_'+col,[w['選單'][f'{col}{i}'].value for i in range(1,len(expected)+1)]==expected)
  check('no_auto_approval_dropdown',all(not any(word in v for word in ['已核准','發布','可用']) for v in STATUSES+DECISIONS+PROTECTED_DECISIONS))
  for id,row in byid.items():
   n=row[0].row;decision=[r for r in rules if f'P{n}' in r.sqref];check(f'{id}_correct_decision_rule',len(decision)==1 and decision[0].formula1==('=ProtectedDecisions' if id in protected else '=ReviewDecisions'))
   for col in ['Q','R']:
    numeric=[r for r in rules if r.type=='decimal' and f'{col}{n}' in r.sqref];check(f'{id}_numeric_rule_and_protection_agree',len(numeric)==(0 if id in protected else 1))
  formulas=[]
  for ws in w:
   for row in ws:
    for cell in row:
     if cell.data_type=='f':
      check(f'formula_internal_{ws.title}_{cell.coordinate}',ws.title in ['高風險優先審核','審核統計'] and '[' not in cell.value and 'HYPERLINK' not in cell.value and ('審核表' in cell.value));formulas.append(f'{ws.title}!{cell.coordinate}')
  check('562_known_internal_formulas',len(formulas)==62*9+4)
  check('auto_recalculation_enabled',w.calculation.calcMode=='auto' and w.calculation.fullCalcOnLoad)
  check('history_rows_present',history.max_row>=636)
  if source:
   check('source_hash_matches_receipt',sha(source)==receipt['sourceSha256'])
   data=load_source(source,allow_ods=True);extra=json.loads(Path(supplement).read_text()) if supplement else None;expected,_=prepare(data,extra)
   for r in expected:
    actual=byid[r['id']]
    for i in list(range(14))+list(range(23,28)):
     a=actual[i].value;e=r['values'][i]
     if a in (None,'') and e in (None,''):ok=True
     elif isinstance(a,(float,int)) and isinstance(e,(float,int)):ok=abs(a-e)<1e-12
     else:ok=a==e
     check(f"{r['id']}_source_alignment_{i+1}",ok)
  with zipfile.ZipFile(path) as z:check('no_macros_external_workbook_links_or_data_connections',not any('vbaProject' in s or 'externalLinks/' in s or 'connections.xml' in s for s in z.namelist()))
 finally:w.close()
 check('review_workbook_unchanged_by_validation',sha(path)==before)
 return {'passed':True,'assertionsPassed':len(checks),'rows':635,'uniqueIds':635,'protectedCount':9,'validCoordinateLinks':valid_links,'invalidCandidateIds':invalid,'nameLinks':635,'highRiskUniqueIds':len(cids),'duplicateGroups':len(groupids),'duplicateUniqueIds':len(dups),'formulaCount':len(formulas),'formulaEvaluationVerified':False,'excelUiVerified':False,'sourcePrimaryXlsxVerified':receipt['requestedXlsxProvided'],'checks':checks}

if __name__=='__main__':
 p=argparse.ArgumentParser(description=__doc__);p.add_argument('workbook',type=Path);p.add_argument('receipt',type=Path);p.add_argument('--source',type=Path);p.add_argument('--supplement',type=Path);p.add_argument('--output',required=True,type=Path);a=p.parse_args();r=validate(a.workbook,json.loads(a.receipt.read_text()),a.source,a.supplement)
 with a.output.open('x') as f:json.dump(r,f,ensure_ascii=False,indent=2);f.write('\n')
 print(json.dumps({k:v for k,v in r.items() if k!='checks'},ensure_ascii=False,indent=2))
