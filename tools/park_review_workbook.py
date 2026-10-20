"""Offline manual-review workbook. No source writes, approvals, sync, or network calls."""
import argparse,collections,csv,datetime as dt,hashlib,json,math,re,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
from urllib.parse import urlencode,urlsplit
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment,Protection
from openpyxl.comments import Comment
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter as letter

PRIORITY=['NP-12','NP-112','NP-333','NP-480','NP-495','NP-541','NP-551','NP-595']
SHEETS={'parks':'公園資料','candidates':'公園座標候選','reviews':'公園座標品質檢核'}
HEADERS=['公園ID','名稱','行政區','地址','正式緯度','正式經度','候選緯度','候選經度','原品質分級','補充品質分級','異常原因','優先順序','候選座標Google Maps連結','公園名稱Google Maps搜尋連結','審核狀態','審核決定','修正緯度','修正經度','座標代表類型','證據URL','審核日期','審核人','備註','正式座標保護','重複座標群組','來源正式列','來源候選列','來源品質列']
STATUSES=['待審','審核中','待補證據','已填寫待複核']
DECISIONS=['採用候選（待複核）','人工修正（待複核）','保留正式座標','排除候選','無法判定']
PROTECTED_DECISIONS=['保留正式座標','待補證據','無法判定']
TYPES=['園區代表點','入口','未知']

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def text(v):return str(v if v is not None else '').strip()
def valid(lat,lon):
 try:
  if isinstance(lat,bool) or isinstance(lon,bool) or text(lat)=='' or text(lon)=='':return False
  return 24.7<=float(lat)<=25.4 and 121.2<=float(lon)<=122.1
 except (TypeError,ValueError):return False

def row_objects(headers,rows,name):
 headers=[text(h) for h in headers]
 while headers and not headers[-1]:headers.pop()
 if not headers or any(not h for h in headers) or len(set(headers))!=len(headers) or '__sheetRow' in headers:raise ValueError(name+': invalid headers')
 out=[]
 for number,values in rows:
  if not any(v not in (None,'') for v in values):continue
  if any(v not in (None,'') for v in values[len(headers):]):raise ValueError(name+': unnamed data column')
  out.append(dict({h:values[i] if i<len(values) and values[i] is not None else '' for i,h in enumerate(headers)},__sheetRow=number))
 return out

def load_source(source,allow_ods=False):
 source=Path(source);before=sha(source)
 if source.stat().st_size>64*1024*1024:raise ValueError('Source too large')
 with zipfile.ZipFile(source) as archive:
  if sum(i.file_size for i in archive.infolist())>256*1024*1024:raise ValueError('Unpacked source too large')
 tables={};formulas=[]
 if source.suffix.lower()=='.xlsx':
  vb=openpyxl.load_workbook(source,read_only=True,data_only=True,keep_links=False);fb=openpyxl.load_workbook(source,read_only=True,data_only=False,keep_links=False)
  try:
   for name in vb.sheetnames:
    if name not in SHEETS.values() and not re.search('B4|人工|修正',name):continue
    vs,fs=vb[name],fb[name]
    if vs.max_row>100000 or vs.max_column>256:raise ValueError('Source dimensions too large')
    values=list(vs.values);raw=list(fs.iter_rows())
    if not values:continue
    if name not in SHEETS.values() and '公園ID' not in values[0]:continue
    tables[name]=row_objects(values[0],enumerate(values[1:],2),name)
    for row in raw:
     for cell in row:
      if cell.data_type=='f':formulas.append({'sheet':name,'cell':cell.coordinate,'cachedValuePresent':values[cell.row-1][cell.column-1] is not None})
  finally:vb.close();fb.close()
  kind='XLSX_PRIMARY_INPUT'
 elif source.suffix.lower()=='.ods' and allow_ods:
  T='{urn:oasis:names:tc:opendocument:xmlns:table:1.0}';O='{urn:oasis:names:tc:opendocument:xmlns:office:1.0}';X='{urn:oasis:names:tc:opendocument:xmlns:text:1.0}'
  with zipfile.ZipFile(source) as z:root=ET.fromstring(z.read('content.xml'))
  for tab in root.findall('.//'+O+'spreadsheet/'+T+'table'):
   name=tab.get(T+'name')
   if name not in SHEETS.values() and not re.search('B4|人工|修正',name):continue
   rows=[];rownum=1;header=None
   for row in tab.findall(T+'table-row'):
    repeat=int(row.get(T+'number-rows-repeated','1'));values=[];formula_flags=[]
    for cell in row:
     if cell.tag not in [T+'table-cell',T+'covered-table-cell']:continue
     count=int(cell.get(T+'number-columns-repeated','1'));typ=cell.get(O+'value-type');formula=cell.get(T+'formula')
     value=cell.get(O+'string-value','\n'.join(''.join(p.itertext()) for p in cell.findall(X+'p')))
     if typ in ['float','currency','percentage']:value=float(cell.get(O+'value'))
     elif typ=='boolean':value=cell.get(O+'boolean-value')=='true'
     elif typ=='date':value=cell.get(O+'date-value','')
     elif typ=='time':value=cell.get(O+'time-value','')
     if value=='' and not formula and count>256:break
     if len(values)+count>256:raise ValueError('ODS columns exceed limit')
     values.extend([value]*count);formula_flags.extend([formula]*count)
    if not any(v!='' for v in values) and not any(formula_flags):rownum+=repeat;continue
    if repeat>100000 or rownum>100000:raise ValueError('ODS rows exceed limit')
    if header is None:header=values;rownum+=repeat;continue
    for k in range(repeat):
     rows.append((rownum+k,values[:]))
     for i,f in enumerate(formula_flags):
      if f:formulas.append({'sheet':name,'cell':f'{letter(i+1)}{rownum+k}','cachedValuePresent':values[i]!=''})
    rownum+=repeat
   if header and (name in SHEETS.values() or '公園ID' in header):tables[name]=row_objects(header,rows,name)
  kind='ODS_DRAFT_REQUESTED_XLSX_MISSING'
 else:raise ValueError('Expected XLSX; ODS requires explicit --allow-ods-draft')
 if any(n not in tables for n in SHEETS.values()):raise ValueError('Missing required sheets')
 if sha(source)!=before:raise ValueError('Source changed during read')
 return {'tables':tables,'evidence':{'filename':source.name,'sourceSha256':before,'sourceFormat':source.suffix.lower(),'inputStatus':kind,'requestedXlsxProvided':kind=='XLSX_PRIMARY_INPUT','formulaCells':formulas}}

def index(rows,scope):
 found={}
 for r in rows:
  id=text(r.get('公園ID'))
  if not id:raise ValueError(scope+': missing ID')
  if id in found:raise ValueError(scope+': duplicate ID '+id)
  found[id]=r
 return found

def prepare(data,supplement=None):
 tables=data['tables'];parks=[p for p in tables[SHEETS['parks']] if re.fullmatch('NP-\\d+',text(p.get('公園ID')))]
 if any(text(p.get('縣市'))!='新北市' for p in parks):raise ValueError('NP ID outside New Taipei')
 P=index(parks,'parks');expected={f'NP-{i}' for i in range(1,636)}
 if set(P)!=expected:raise ValueError('Exactly NP-1 through NP-635 required; no drops or substitutions')
 indexes={}
 for key in ['candidates','reviews']:
  rows=tables[SHEETS[key]]
  if any(not text(r.get('公園ID')) for r in rows):raise ValueError(key+': missing ID')
  selected=[r for r in rows if text(r.get('公園ID')).startswith('NP-')];indexes[key]=index(selected,key)
  if set(indexes[key])!=expected:raise ValueError(key+': ID set mismatch')
 C=indexes['candidates'];R=indexes['reviews']
 for id in P:
  if any(text(r.get('公園名稱'))!=text(P[id].get('公園名稱')) for r in [C[id],R[id]]):raise ValueError(id+': names disagree between sheets')
 groups=collections.defaultdict(set)
 for id in sorted(P,key=lambda x:int(x[3:])):
  for lat,lon in [(C[id].get('候選緯度'),C[id].get('候選經度')),(P[id].get('緯度'),P[id].get('經度'))]:
   if valid(lat,lon):groups[f'{float(lat):.6f},{float(lon):.6f}'].add(id)
 groups=[{'groupId':f'G{i:02}','coordinate':coordinate,'ids':sorted(ids,key=lambda x:int(x[3:]))} for i,(coordinate,ids) in enumerate(((k,v) for k,v in groups.items() if len(v)>1),1)]
 memberships=collections.defaultdict(list)
 for group in groups:
  for id in group['ids']:memberships[id].append(group['groupId'])
 prior={r['id']:r for r in supplement} if supplement is not None else None
 if prior is not None:
  if set(prior)!=expected:raise ValueError('Supplement ID set mismatch')
  for id,p in P.items():
   q=prior[id];c=C[id]
   if text(q.get('name'))!=text(p.get('公園名稱')):raise ValueError('Stale supplement name '+id)
   if text(q.get('district'))!=text(p.get('行政區')) or text(q.get('officialAddress'))!=text(p.get('地址')):raise ValueError('Stale supplement address '+id)
   if text(q.get('googleAddress'))!=text(c.get('Google回傳地址')) or text(q.get('locationType'))!=text(c.get('定位類型')):raise ValueError('Stale supplement geocoding evidence '+id)
   for a,b in [(p.get('緯度'),q.get('officialLatitude')),(p.get('經度'),q.get('officialLongitude')),(c.get('候選緯度'),q.get('candidateLatitude')),(c.get('候選經度'),q.get('candidateLongitude'))]:
    if text(a)==text(b):continue
    try:equal=abs(float(a)-float(b))<1e-12
    except (TypeError,ValueError):equal=False
    if not equal:raise ValueError('Stale supplement coordinate '+id)
 records=[]
 for id,p in P.items():
  c=C[id];r=R[id];protected=valid(p.get('緯度'),p.get('經度'));cv=valid(c.get('候選緯度'),c.get('候選經度'))
  if prior is not None:
   grade=prior[id]['grade'];reasons=prior[id]['severeReasons']+prior[id]['warnings']
  else:
   severe=[];warnings=[];fmt=text(c.get('Google回傳地址')).replace('芦','蘆').replace('区','區');district=text(p.get('行政區'))
   if not cv:severe.append('候選座標無效或缺漏')
   if id in memberships:severe.append('不同公園共用座標')
   if any(v in fmt for v in ['高雄市','桃園市','基隆市','台北市','臺北市']):severe.append('Google明示其他縣市')
   if any(text(other.get('行政區')) in fmt for other in P.values() if text(other.get('行政區')) and text(other.get('行政區'))!=district):severe.append('Google明示其他行政區')
   if fmt and district not in fmt:warnings.append('Google地址與行政區需人工比對')
   if protected and cv:
    a,b,x,y=map(math.radians,map(float,[p['緯度'],p['經度'],c['候選緯度'],c['候選經度']]))
    delta=6371000*2*math.asin(min(1,math.sqrt(math.sin((x-a)/2)**2+math.cos(a)*math.cos(x)*math.sin((y-b)/2)**2)))
    if delta>1000:severe.append('候選與正式座標相距超過1km')
    elif delta>100:warnings.append('候選與正式座標相距超過100m')
   if c.get('定位類型')!='ROOFTOP':warnings.append('定位類型不是入口證據')
   addr=text(p.get('地址')).removeprefix('新北市').removeprefix(district)
   if not addr or addr not in fmt:warnings.append('地址一致性證據不足')
   grade='C' if severe else 'B' if warnings else 'A';reasons=severe+warnings
  if grade not in ['A','B','C']:raise ValueError('Invalid supplementary grade')
  if not cv and '候選座標無效或缺漏' not in reasons:reasons+=['候選座標無效或缺漏']
  if protected:reasons+=['既有正式座標保護；不得在本工具覆蓋']
  if id in memberships:reasons+=['共點組：'+'、'.join(memberships[id])+'；不代表同一公園，勿合併ID']
  url='https://www.google.com/maps/search/?'+urlencode({'api':'1','query':f"{c['候選緯度']},{c['候選經度']}"}) if cv else ''
  nameurl='https://www.google.com/maps/search/?'+urlencode({'api':'1','query':f"{p['公園名稱']} 新北市 {p['行政區']}"})
  priority='P0｜指定優先' if id in PRIORITY else 'P1｜C級高風險' if grade=='C' else 'P1｜正式保護' if protected else 'P2｜一般人工審核'
  values=[id,p['公園名稱'],p['行政區'],p['地址'],p.get('緯度',''),p.get('經度',''),c.get('候選緯度',''),c.get('候選經度',''),r.get('檢核分級',''),grade,'；'.join(dict.fromkeys(reasons)) or '仍須人工比對園區／入口',priority,url,nameurl,'待審','','','','未知','','','','','保護' if protected else '','、'.join(memberships[id]),p['__sheetRow'],c['__sheetRow'],r['__sheetRow']]
  records.append({'id':id,'values':values,'protected':protected,'grade':grade,'candidateValid':cv,'groups':memberships[id]})
 def order(r):return (0,PRIORITY.index(r['id'])) if r['id'] in PRIORITY else (1 if r['grade']=='C' else 2 if r['protected'] else 3 if r['grade']=='B' else 4,int(r['id'][3:]))
 records.sort(key=order)
 return records,groups

NAVY='17365D';GREEN='E2F0D9';GREY='E7E6E6';YELLOW='FFF2CC';RED='FCE4D6'
def literal(cell,value):
 if isinstance(value,str) and len(value)>32767:raise ValueError('Cell text exceeds Excel limit')
 cell.value=value
 if isinstance(value,str):cell.data_type='s'
 cell.alignment=Alignment(vertical='top',wrap_text=True)
def header(ws,cols):
 ws.append(cols);ws.freeze_panes='C2';ws.auto_filter.ref=f'A1:{letter(len(cols))}{max(1,ws.max_row)}';ws.sheet_view.zoomScale=80
 for c in ws[1]:c.font=Font(name='Calibri',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor=NAVY);c.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=42

def make_workbook(data,records,groups,output):
 wb=openpyxl.Workbook();intro=wb.active;intro.title='使用說明';main=wb.create_sheet('審核表');stats=wb.create_sheet('審核統計');risk=wb.create_sheet('高風險優先審核');history=wb.create_sheet('既有人工修正紀錄');lists=wb.create_sheet('選單');lists.sheet_state='hidden';wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
 wb.properties.title='PapaKing 新北635公園人工審核';wb.properties.subject=data['evidence']['inputStatus'];wb.properties.creator='PapaKing offline review generator'
 descriptions=[('PapaKing｜新北635公園人工審核','只供人工填寫；不會同步、核准或發布任何座標。'),('來源狀態',data['evidence']['inputStatus']),('來源原檔',data['evidence']['filename']),('SHA-256',data['evidence']['sourceSha256']),('重要限制','未取得指定XLSX；本版依既有ODS製作，需先核對是否為最新版本。' if not data['evidence']['requestedXlsxProvided'] else '本版使用指定XLSX；不代表資料已核准或來源為最新。'),('開始審核','開啟「審核表」；前8筆為指定優先，接著為C級。用標題篩選；ID及來源欄不能更改。'),('顏色','藍灰＝鎖定來源；綠色＝人工輸入；灰色修正座標＝正式資料保護，不能填寫。'),('Google Maps','座標連結只供位置比對，不是核准或入口證據；無效候選僅提供名稱搜尋。'),('人工步驟','先選審核中，核對園區或入口及來源，再選決定；必要時填修正座標、URL、日期與審核人。'),('保護規則','9筆有效正式座標標示保護；本輪不接受覆蓋，疑義請記備註另案處理。'),('共點規則','同點組可能重疊；不能把同點視為同園、不能合併ID。統計均按唯一ID計數。'),('完成後','選已填寫待複核；此狀態不等於核准。黃／紅色缺漏提示需補齊。'),('高風險頁','只供查閱，點公園ID跳回主表填寫，避免兩份輸入不同步。'),('歷史紀錄','原有已確認／需要修正另存唯讀歷史頁；本輪狀態仍全部待審，輸入欄不預填。'),('儲存與交回','另存副本如PapaKing_Review_姓名_日期.xlsx，交回完整.xlsx，勿只交CSV或刪除列／改ID。'),('公式','統計與高風險頁使用內部公式；Excel開啟時計算。openpyxl不計算快取，不能宣稱已驗證Excel顯示結果。'),('安全','鎖定僅防誤改，不是加密或存取控制。不要解除保護；本工具沒有任何正式寫入功能。')]
 header(intro,['項目','操作／說明'])
 for row in descriptions:
  intro.append(row);intro.row_dimensions[intro.max_row].height=42
 intro.column_dimensions['A'].width=24;intro.column_dimensions['B'].width=115;intro.freeze_panes='A2'
 header(main,HEADERS)
 for r in records:
  main.append([None]*len(HEADERS));num=main.max_row
  for i,v in enumerate(r['values'],1):
   c=main.cell(num,i);literal(c,v);c.font=Font(name='Calibri',size=11);c.fill=PatternFill('solid',fgColor='F2F5F9' if num%2==0 else 'FFFFFF')
   if 15<=i<=23:c.protection=Protection(locked=False);c.fill=PatternFill('solid',fgColor=GREEN)
   if r['protected'] and i in [17,18]:c.protection=Protection(locked=True);c.fill=PatternFill('solid',fgColor=GREY);c.comment=Comment('已有正式座標，禁止本輪修正；請記錄疑義另案處理。','PapaKing')
   if i in [13,14] and v:c.hyperlink=v;c.font=Font(name='Calibri',size=10,color='0563C1',underline='single')
   if i in [5,6,7,8,17,18]:c.number_format='0.0000000'
   if i==21:c.number_format='yyyy-mm-dd'
  main.row_dimensions[num].height=70
 end=main.max_row;main.auto_filter.ref=f'A1:AB{end}';main.print_title_rows='1:1';main.protection.sheet=True;main.protection.autoFilter=False;main.protection.sort=True;main.protection.selectLockedCells=False;main.protection.selectUnlockedCells=False
 widths=[13,29,12,42,16,16,16,16,21,15,65,23,34,34,22,27,16,16,24,40,15,16,40,17,19,15,15,15]
 for i,w in enumerate(widths,1):main.column_dimensions[letter(i)].width=w
 for col in ['O','P','Q','R','S','T','U','V','W']:main[f'{col}1'].comment=Comment('人工輸入欄；未經複核，不代表核准。','PapaKing')
 for col,values in [('A',STATUSES),('B',DECISIONS),('C',TYPES),('D',PROTECTED_DECISIONS)]:
  for i,value in enumerate(values,1):literal(lists[f'{col}{i}'],value)
 for name,col,values in [('ReviewStatuses','A',STATUSES),('ReviewDecisions','B',DECISIONS),('PositionTypes','C',TYPES),('ProtectedDecisions','D',PROTECTED_DECISIONS)]:wb.defined_names.add(DefinedName(name,attr_text=f"'選單'!${col}$1:${col}${len(values)}"))
 def dv(type,formula1,formula2=None):
  rule=DataValidation(type=type,formula1=formula1,formula2=formula2,allow_blank=True);rule.errorTitle='請檢查輸入';rule.error='請使用下拉選項或有效數值；未知請留空。';rule.showErrorMessage=True;rule.errorStyle='stop';rule.showInputMessage=True;rule.promptTitle='人工審核';rule.prompt='這些輸入不會自動核准或同步。';main.add_data_validation(rule);return rule
 dv('list','=ReviewStatuses').add(f'O2:O{end}');dv('list','=PositionTypes').add(f'S2:S{end}')
 regular=dv('list','=ReviewDecisions');protected=dv('list','=ProtectedDecisions');lat=dv('decimal','24.7','25.4');lat.operator='between';lon=dv('decimal','121.2','122.1');lon.operator='between'
 for n,r in enumerate(records,2):
  (protected if r['protected'] else regular).add(f'P{n}')
  if not r['protected']:lat.add(f'Q{n}');lon.add(f'R{n}')
 date=dv('date','DATE(2000,1,1)','DATE(2100,12,31)');date.operator='between';date.add(f'U2:U{end}')
 dv('custom','OR(T2="",AND(OR(LEFT(T2,8)="https://",LEFT(T2,7)="http://"),ISERROR(SEARCH(" ",T2))))').add(f'T2:T{end}')
 main.conditional_formatting.add(f'A2:N{end}',FormulaRule(formula=['$J2="C"'],fill=PatternFill('solid',fgColor=RED)))
 main.conditional_formatting.add(f'X2:X{end}',FormulaRule(formula=['$X2="保護"'],fill=PatternFill('solid',fgColor=YELLOW)))
 main.conditional_formatting.add(f'Q2:R{end}',FormulaRule(formula=['COUNTA($Q2:$R2)=1'],fill=PatternFill('solid',fgColor=RED)))
 main.conditional_formatting.add(f'O2:W{end}',FormulaRule(formula=['AND($O2="已填寫待複核",OR($P2="",$T2="",$U2="",$V2="",AND($P2="人工修正（待複核）",COUNT($Q2:$R2)<>2)))'],fill=PatternFill('solid',fgColor=YELLOW)))
 header(risk,HEADERS);locations={r['id']:i for i,r in enumerate(records,2)}
 for r in [r for r in records if r['grade']=='C']:
  risk.append([None]*len(HEADERS));n=risk.max_row;master=locations[r['id']]
  for i,v in enumerate(r['values'],1):
   c=risk.cell(n,i)
   if 15<=i<=23:c.value=f'=IF(\'審核表\'!{letter(i)}{master}="","",\'審核表\'!{letter(i)}{master})'
   else:literal(c,v)
   c.alignment=Alignment(wrap_text=True,vertical='top')
   if i in [13,14] and v:c.hyperlink=v;c.font=Font(color='0563C1',underline='single')
   if i in [5,6,7,8,17,18]:c.number_format='0.0000000'
   if i==21:c.number_format='yyyy-mm-dd'
  risk.cell(n,1).hyperlink=Hyperlink(ref=f'A{n}',location=f"'審核表'!A{master}",display=r['id']);risk.cell(n,1).font=Font(color='0563C1',underline='single');risk.row_dimensions[n].height=65
 for i,w in enumerate(widths,1):risk.column_dimensions[letter(i)].width=w
 risk.auto_filter.ref=f'A1:AB{risk.max_row}';risk.protection.sheet=True;risk.protection.autoFilter=False
 histcols=['來源分頁','來源列','公園ID','原人工審核結果','原修正緯度','原修正經度','原修正座標來源','原修正座標驗收','原人工審核備註','原參考連結','原檢核原因','其他原始紀錄JSON'];header(history,histcols)
 ids={r['id'] for r in records}
 for name,rows in data['tables'].items():
  if name in [SHEETS['parks'],SHEETS['candidates']]:continue
  for r in rows:
   if text(r.get('公園ID')) not in ids:continue
   values=[name,r['__sheetRow'],r['公園ID'],r.get('人工審核結果',''),r.get('修正後緯度',r.get('修正緯度','')),r.get('修正後經度',r.get('修正經度','')),r.get('修正座標來源',r.get('來源','')),r.get('修正座標驗收',''),r.get('人工審核備註',r.get('備註','')),r.get('Google 地圖參考連結',r.get('Google 地圖查詢','')),r.get('檢核原因',''),json.dumps(r,ensure_ascii=False,default=str)]
   history.append([None]*len(values))
   for i,v in enumerate(values,1):literal(history.cell(history.max_row,i),v)
 for i,w in enumerate([25,12,13,22,17,17,70,18,80,40,60,90],1):history.column_dimensions[letter(i)].width=w
 history.auto_filter.ref=f'A1:L{history.max_row}';history.protection.sheet=True;history.protection.autoFilter=False
 header(stats,['統計項目','數值／狀態'])
 fixed=[('來源狀態',data['evidence']['inputStatus']),('唯一公園ID',len(ids)),('正式保護ID',sum(r['protected'] for r in records)),('高風險C級唯一ID',sum(r['grade']=='C' for r in records)),('重複座標群組',len(groups)),('共點唯一ID',len({id for g in groups for id in g['ids']})),('指定八筆與共點交集',len(set(PRIORITY)&{id for g in groups for id in g['ids']})),('有效候選',sum(r['candidateValid'] for r in records)),('無效候選',sum(not r['candidateValid'] for r in records)),('新增核准／發布',0)]
 for r in fixed:stats.append(r)
 for status in STATUSES:
  n=stats.max_row+1;stats.append([status,f'=COUNTIF(\'審核表\'!$O$2:$O${end},A{n})'])
 stats.column_dimensions['A'].width=35;stats.column_dimensions['B'].width=65;stats.freeze_panes='A2';stats.protection.sheet=True
 for ws in [intro,stats]:
  for row in ws.iter_rows(min_row=2):
   for c in row:
    if c.data_type!='f':literal(c,c.value)
    c.alignment=Alignment(wrap_text=True,vertical='top')
 if output.exists():raise FileExistsError(output)
 wb.save(output)
 return wb

def build(source,out,allow_ods=False,supplement_path=None):
 data=load_source(source,allow_ods);supplement=json.loads(Path(supplement_path).read_text()) if supplement_path else None
 records,groups=prepare(data,supplement)
 if {r['id'] for r in records if r['protected']}!=set(PRIORITY+['NP-9']):raise ValueError('Expected nine known formal coordinates; investigate source mismatch before build')
 out=Path(out);out.mkdir(mode=0o700)
 workbook=out/'PapaKing_NTPC_635_Park_Coordinate_Review.xlsx';make_workbook(data,records,groups,workbook)
 def csvout(filename,rows):
  with (out/filename).open('x',encoding='utf-8-sig',newline='') as f:
   w=csv.writer(f);w.writerow(HEADERS)
   for r in rows:
    values=["'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')) else v for v in r['values']];w.writerow(values)
 csvout('PapaKing_NTPC_635_Review_Base.csv',records);csvout('PapaKing_NTPC_High_Risk_Review.csv',[r for r in records if r['grade']=='C'])
 receipt={**data['evidence'],'outputWorkbook':str(workbook.resolve()),'parkCount':len(records),'uniqueIds':len({r['id'] for r in records}),'protectedIds':[r['id'] for r in records if r['protected']],'validCandidates':sum(r['candidateValid'] for r in records),'invalidCandidateIds':[r['id'] for r in records if not r['candidateValid']],'highRiskCount':sum(r['grade']=='C' for r in records),'grades':dict(collections.Counter(r['grade'] for r in records)),'duplicateGroups':groups,'duplicateUniqueIds':len({id for g in groups for id in g['ids']}),'supplementUsed':supplement_path is not None,'supplementSha256':sha(supplement_path) if supplement_path else None,'sourceUnchanged':sha(source)==data['evidence']['sourceSha256'],'newApprovals':0,'sourceWrites':0,'formulaCalculation':'not evaluated by openpyxl; auto recalculation on Excel open'}
 (out/'source-and-build-receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2,default=str)+'\n')
 (out/'README.md').write_text(f'''# PapaKing 新北635公園人工審核

來源：{data['evidence']['filename']}；SHA-256：{data['evidence']['sourceSha256']}。
來源狀態：**{data['evidence']['inputStatus']}**。
{'未取得本輪指定XLSX。此檔以先前ODS建立草案，使用前須核對最新XLSX。' if not data['evidence']['requestedXlsxProvided'] else '本輪讀取XLSX主輸入。'}

開啟.xlsx使用「使用說明」「審核表」「高風險優先審核」「審核統計」「既有人工修正紀錄」。所有635筆預設待審；決定、修正座標、證據URL、日期、審核人及備註留空；座標代表類型預設未知。

先看前8筆，再按高風險頁或主表篩選。來源欄鎖定；綠色欄供人工輸入。9筆正式座標標記保護，修正座標欄鎖定。高風險頁只供查閱，點ID回主表填寫。共點不等於同園，不得合併或刪除ID。

點候選座標連結或公園名稱搜尋，核對官方園區或入口，保存來源URL、審核日期與姓名。無效候選不產生座標連結。Google搜尋結果、A級或官方名稱相符均不是核准。選已填寫待複核前補齊必要證據；黃色提示需補齊，紅色提示座標只填一欄。

另存副本，例如PapaKing_Review_姓名_20261020.xlsx，交回完整.xlsx。勿只交CSV、勿改ID、勿刪列、勿解除來源保護；疑義記備註。CSV是基礎資料，不含Excel下拉或保護功能。

公式只在統計與高風險頁，openpyxl未計算其顯示快取，Excel開啟會重算。工作表保護僅防誤操作，不是加密、權限或惡意修改防護。回收結果仍須重新驗ID、來源指紋、證據及人工複核。

本工具不連網、不寫Google、不執行B4-5、不核准或發布座標；歷史已確認內容只放唯讀歷史頁，不預填本輪決定。
''',encoding='utf-8')
 return receipt

if __name__=='__main__':
 p=argparse.ArgumentParser(description=__doc__);p.add_argument('source',type=Path);p.add_argument('output_dir',type=Path);p.add_argument('--allow-ods-draft',action='store_true');p.add_argument('--supplement',type=Path);a=p.parse_args();print(json.dumps(build(a.source,a.output_dir,a.allow_ods_draft,a.supplement),ensure_ascii=False,indent=2))
