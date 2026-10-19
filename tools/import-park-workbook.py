#!/usr/bin/env python3
"""Read four sheets from a user-exported XLSX; never authenticate or change source.
Requires openpyxl (validated version recorded in the test log). No network calls.
"""
import argparse
import datetime as dt
import hashlib
import json
import os
from pathlib import Path
import zipfile
import openpyxl

SHEETS = {'parks': '公園資料', 'candidates': '公園座標候選', 'facilities': '公園設施', 'reviews': '公園座標品質檢核'}
MAX_SOURCE_BYTES = 64 * 1024 * 1024
MAX_UNPACKED_BYTES = 256 * 1024 * 1024
MAX_ROWS = 100000
MAX_COLUMNS = 256


def value_json(value):
    if value is None:
        return ''
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return value


def import_workbook(source, mapping=None):
    source = Path(source)
    if source.stat().st_size > MAX_SOURCE_BYTES:
        raise ValueError('XLSX exceeds 64 MiB; provide a four-sheet local copy')
    source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
    with zipfile.ZipFile(source) as archive:
        if sum(info.file_size for info in archive.infolist()) > MAX_UNPACKED_BYTES:
            raise ValueError('Unpacked XLSX exceeds limit')
    values_book = openpyxl.load_workbook(source, read_only=True, data_only=True, keep_links=False)
    formula_book = openpyxl.load_workbook(source, read_only=True, data_only=False, keep_links=False)
    mapping = mapping or {}
    try:
        missing = [name for name in SHEETS.values() if name not in values_book.sheetnames]
        if missing:
            raise ValueError('Missing required sheets: ' + '、'.join(missing))
        if any(key not in SHEETS for key in mapping):
            raise ValueError('Unknown mapping table')
        result = {}
        inventory = {}
        for key, name in SHEETS.items():
            vs, fs = values_book[name], formula_book[name]
            if vs.max_row > MAX_ROWS or vs.max_column > MAX_COLUMNS:
                raise ValueError(name + ': dimensions exceed safe import limit')
            headers = list(next(fs.iter_rows(min_row=1, max_row=1, values_only=True)))
            while headers and headers[-1] is None:
                headers.pop()
            if not headers or any(not isinstance(h, str) or not h.strip() or h.startswith('=') for h in headers):
                raise ValueError(name + ': missing, blank or non-text header')
            headers = [h.strip() for h in headers]
            if len(set(headers)) != len(headers) or '__sheetRow' in headers:
                raise ValueError(name + ': duplicate or reserved header')
            aliases = mapping.get(key, {})
            for canonical, actual in aliases.items():
                if actual not in headers:
                    raise ValueError(name + ': mapping source absent: ' + actual)
                if canonical in headers and canonical != actual:
                    raise ValueError(name + ': mapping would overwrite original column: ' + canonical)
            rows, formula_cells = [], []
            for row_number, (vr, fr) in enumerate(zip(vs.iter_rows(min_row=2), fs.iter_rows(min_row=2), strict=True), 2):
                if not any(c.value is not None for c in fr):
                    continue
                if any(c.value is not None for c in fr[len(headers):]):
                    raise ValueError(name + ': data outside named columns at row ' + str(row_number))
                obj = {h: value_json(vr[i].value) for i, h in enumerate(headers)}
                obj['__sheetRow'] = row_number
                for i, cell in enumerate(fr[:len(headers)]):
                    if cell.data_type == 'f':
                        formula_cells.append({'row': row_number, 'column': headers[i], 'cachedValuePresent': vr[i].value is not None})
                for canonical, actual in aliases.items():
                    obj[canonical] = obj[actual]
                rows.append(obj)
            result[key] = rows
            inventory[key] = {'sheet': name, 'headers': headers, 'dataRows': len(rows),
                              'formulaCells': formula_cells, 'columnMapping': aliases}
        if hashlib.sha256(source.read_bytes()).hexdigest() != source_sha256:
            raise ValueError('Source XLSX changed during import; preserve a stable export and rerun')
        result['exportEvidence'] = {
            'format': 'papaking-workbook-export-v1',
            'sourceSha256': source_sha256,
            'importedAt': dt.datetime.now(dt.timezone.utc).isoformat(),
            'googleExportedAt': None, 'googleRevision': None,
            'dateEncoding': 'XLSX dates converted to ISO strings without inventing a timezone',
            'source': 'User downloaded Google workbook; not an authenticated API read or Sheets write verification',
            'inventory': inventory,
            'extraSheetsExcluded': len(values_book.sheetnames) - len(SHEETS),
        }
        return result
    finally:
        values_book.close()
        formula_book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input_xlsx')
    parser.add_argument('new_output_json')
    parser.add_argument('--mapping', help='Explicit canonical-to-source header mapping, grouped by parks/candidates/facilities/reviews')
    args = parser.parse_args()
    mapping = json.loads(Path(args.mapping).read_text()) if args.mapping else None
    result = import_workbook(args.input_xlsx, mapping)
    # Exclusive creation protects the export and earlier evidence; outputs stay private.
    fd = os.open(args.new_output_json, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(fd, 'w', encoding='utf-8') as stream:
        json.dump(result, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write('\n')
    print(json.dumps({k: v['dataRows'] for k, v in result['exportEvidence']['inventory'].items()}, ensure_ascii=False))


if __name__ == '__main__':
    main()
