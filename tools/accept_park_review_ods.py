"""Reconcile an existing PK-02E workbook with the user-designated ODS primary.
Never overwrites the input workbook or source. Only cover/source labels are changed
in a new ZIP copy after all source cells and historical records match.
"""
import argparse
import csv
import datetime as dt
import json
import posixpath
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl
from park_review_workbook import HEADERS, load_source, make_workbook, prepare, sha

STATUS = 'ODS_PRIMARY_INPUT_ACCEPTED'
NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def equal(a, b):
    if a in (None, '') and b in (None, ''):
        return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) < 1e-12
    return a == b


def reconcile(source, workbook, supplement=None):
    """Read-only reconciliation; human inputs O:W deliberately excluded."""
    source, workbook = Path(source), Path(workbook)
    if source.suffix.lower() != '.ods':
        raise ValueError('This acceptance workflow requires the designated ODS')
    hashes = {'sourceSha256': sha(source), 'inputWorkbookSha256': sha(workbook)}
    data = load_source(source, allow_ods=True)
    extra = json.loads(Path(supplement).read_text()) if supplement else None
    records, groups = prepare(data, extra)
    diffs = []

    def difference(scope, id, field, expected, actual):
        if not equal(expected, actual):
            diffs.append({'scope': scope, 'parkId': id, 'field': field,
                          'sourceValue': expected, 'workbookValue': actual})

    w = openpyxl.load_workbook(workbook, data_only=False, keep_links=False)
    try:
        main = w['審核表']
        difference('structure', '', 'headers', HEADERS, list(next(main.values)))
        rows = list(main.iter_rows(min_row=2))
        byid = {r[0].value: r for r in rows}
        difference('structure', '', 'rowCount', 635, len(rows))
        difference('structure', '', 'uniqueIdCount', 635, len(byid))
        difference('structure', '', 'idSet', sorted(r['id'] for r in records), sorted(str(i) for i in byid))
        for record in records:
            actual = byid.get(record['id'])
            if actual is None:
                continue
            for i in list(range(14)) + list(range(23, 28)):
                difference('source', record['id'], HEADERS[i], record['values'][i], actual[i].value)
        # Independently reconstruct the read-only quality/manual history, including
        # every original field in its JSON column, rather than grade alone.
        with tempfile.TemporaryDirectory() as tmp:
            expected_path = Path(tmp) / 'expected.xlsx'
            expected = make_workbook(data, records, groups, expected_path)
            expected.close()
            expected = openpyxl.load_workbook(expected_path, data_only=False)
            try:
                hs = w['既有人工修正紀錄']
                eh = expected['既有人工修正紀錄']
                difference('history', '', 'rowCount', eh.max_row, hs.max_row)
                for row in eh:
                    id = eh.cell(row[0].row, 3).value
                    for c in row:
                        difference('history', id, c.coordinate, c.value, hs[c.coordinate].value)
            finally:
                expected.close()
    finally:
        w.close()
    if sha(source) != hashes['sourceSha256'] or sha(workbook) != hashes['inputWorkbookSha256']:
        raise ValueError('An input changed during reconciliation')
    return data, {'utc': dt.datetime.now(dt.timezone.utc).isoformat(), **hashes,
                  'rows': 635, 'differences': diffs, 'differenceCount': len(diffs),
                  'sourcePrimaryOdsVerified': not diffs,
                  'supplementSha256': sha(supplement) if supplement else None,
                  'humanInputComparedAsSource': False}


def relabel_copy(workbook, output, status=STATUS):
    """Patch only inline source-label cells; all other ZIP members byte-identical."""
    updates = {'使用說明': {'B3': status, 'B6': 'ODS為使用者指定主來源；635筆來源及品質歷史核對一致。人工審核與正式發布仍各自需要複核。'},
               '審核統計': {'B2': status}}
    with zipfile.ZipFile(workbook) as z:
        book = ET.fromstring(z.read('xl/workbook.xml'))
        rels = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        targets = {r.attrib['Id']: r.attrib['Target'] for r in rels}
        patches = {}
        for sheet in book.find('m:sheets', NS):
            if sheet.attrib['name'] in updates:
                rid = sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
                target = targets[rid]
                path = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
                root = ET.fromstring(z.read(path))
                for address, value in updates[sheet.attrib['name']].items():
                    cell = root.find(f'.//m:c[@r="{address}"]', NS)
                    if cell is None or cell.attrib.get('t') != 'inlineStr':
                        raise ValueError('Expected inline source-label cell: ' + address)
                    for child in list(cell):
                        cell.remove(child)
                    inline = ET.SubElement(cell, '{' + NS['m'] + '}is')
                    ET.SubElement(inline, '{' + NS['m'] + '}t').text = value
                patches[path] = ET.tostring(root, encoding='utf-8')
        if len(patches) != 2:
            raise ValueError('Expected exactly two source-label sheets')
        core = ET.fromstring(z.read('docProps/core.xml'))
        subject = core.find('{http://purl.org/dc/elements/1.1/}subject')
        if subject is None:
            raise ValueError('Expected source subject metadata')
        subject.text = status
        patches['docProps/core.xml'] = ET.tostring(core, encoding='utf-8')
        with zipfile.ZipFile(output, 'x') as out:
            for member in z.infolist():
                out.writestr(member, patches.get(member.filename, z.read(member.filename)))
    with zipfile.ZipFile(workbook) as old, zipfile.ZipFile(output) as new:
        unchanged = [n for n in old.namelist() if n not in patches]
        if old.namelist() != new.namelist() or not all(old.read(n) == new.read(n) for n in unchanged):
            raise AssertionError('Non-label XML changed')
    return {'changedMembers': sorted(patches), 'byteIdenticalOtherMembers': len(unchanged),
            'allReviewAndHumanInputXmlUnchanged': True}


def accept(source, workbook, output_dir, supplement=None):
    data, report = reconcile(source, workbook, supplement)
    out = Path(output_dir)
    out.mkdir(mode=0o700)  # never overwrite a previous review output
    with (out / 'source-differences.csv').open('x', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=['scope', 'parkId', 'field', 'sourceValue', 'workbookValue'])
        writer.writeheader()
        for diff in report['differences']:
            writer.writerow({k: ("'" + v if isinstance(v, str) and v.lstrip().startswith(('=', '+', '-', '@')) else v) for k, v in diff.items()})
    if report['sourcePrimaryOdsVerified']:
        target = out / Path(workbook).name
        report['preservation'] = relabel_copy(workbook, target)
        report['outputWorkbookSha256'] = sha(target)
        report['inputStatus'] = STATUS
        report['outputWorkbook'] = str(target.resolve())
    else:
        report['inputStatus'] = 'SOURCE_DIFFERENCES_REVIEW_REQUIRED'
    if sha(source) != report['sourceSha256'] or sha(workbook) != report['inputWorkbookSha256']:
        raise AssertionError('Original input changed')
    report.update({'originalInputsUnchanged': True, 'newApprovals': 0,
                   'excelUiVerified': False, 'formulaEvaluationVerified': False})
    (out / 'ods-primary-acceptance.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    return report


if __name__ == '__main__':
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('source', type=Path)
    p.add_argument('workbook', type=Path)
    p.add_argument('output_dir', type=Path)
    p.add_argument('--supplement', type=Path)
    a = p.parse_args()
    result = accept(a.source, a.workbook, a.output_dir, a.supplement)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result['sourcePrimaryOdsVerified'] else 1)
