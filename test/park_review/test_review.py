import unittest,sys,copy,tempfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[2]/'tools'))
import park_review_workbook as review
from openpyxl import load_workbook
class ReviewTests(unittest.TestCase):
 def fixture(self):
  tables={name:[] for name in review.SHEETS.values()}
  for i in range(1,636):
   id=f'NP-{i}';tables['公園資料'].append({'公園ID':id,'縣市':'新北市','行政區':'板橋區','公園名稱':f'合成園{i}','地址':f'甲路{i}號','緯度':25.01 if i==9 else '', '經度':121.5 if i==9 else '', '__sheetRow':i+2})
   tables['公園座標候選'].append({'公園ID':id,'公園名稱':f'合成園{i}','候選緯度':25+i*.0001 if i!=480 else '', '候選經度':121.5 if i!=480 else '', '定位類型':'ROOFTOP','Google回傳地址':f'新北市板橋區甲路{i}號','__sheetRow':i+2})
   tables['公園座標品質檢核'].append({'公園ID':id,'公園名稱':f'合成園{i}','檢核分級':'歷史分級','人工審核結果':'已核准','修正後緯度':25,'修正後經度':121.5,'修正座標驗收':'已確認','__sheetRow':i+2})
  return {'tables':tables,'evidence':{'filename':'synthetic.xlsx','sourceSha256':'a'*64,'inputStatus':'XLSX_PRIMARY_INPUT','requestedXlsxProvided':True}}
 def test_boolean_invalid(self):self.assertFalse(review.valid(True,121.5))
 def test_nan_invalid(self):self.assertFalse(review.valid(float('nan'),121.5))
 def test_outside_ntpc_invalid(self):self.assertFalse(review.valid(22,120.5))
 def test_empty_invalid(self):self.assertFalse(review.valid('',121.5))
 def test_exact635(self):self.assertEqual(len(review.prepare(self.fixture())[0]),635)
 def test_missing_rejected(self):
  d=self.fixture();d['tables']['公園資料'].pop()
  with self.assertRaises(ValueError):review.prepare(d)
 def test_duplicate_rejected(self):
  d=self.fixture();d['tables']['公園座標候選'].append(d['tables']['公園座標候選'][0])
  with self.assertRaises(ValueError):review.prepare(d)
 def test_wrong_name_rejected(self):
  d=self.fixture();d['tables']['公園座標品質檢核'][0]['公園名稱']='不同'
  with self.assertRaises(ValueError):review.prepare(d)
 def test_changed_google_evidence_rejects_stale_grade(self):
  d=self.fixture();candidates={r['公園ID']:r for r in d['tables']['公園座標候選']};prior=[]
  for p in d['tables']['公園資料']:
   c=candidates[p['公園ID']];prior.append({'id':p['公園ID'],'name':p['公園名稱'],'district':p['行政區'],'officialAddress':p['地址'],'officialLatitude':p['緯度'],'officialLongitude':p['經度'],'candidateLatitude':c['候選緯度'],'candidateLongitude':c['候選經度'],'googleAddress':c['Google回傳地址'],'locationType':c['定位類型'],'grade':'A','severeReasons':[],'warnings':[]})
  prior[0]['googleAddress']='old address'
  with self.assertRaises(ValueError):review.prepare(d,prior)
 def test_no_historical_approval_copied(self):
  rows,_=review.prepare(self.fixture());self.assertTrue(all(r['values'][14]=='待審' and all(r['values'][j]=='' for j in [15,16,17,19,20,21,22]) for r in rows))
 def test_invalid_has_no_location_link(self):
  rows,_=review.prepare(self.fixture());r=next(r for r in rows if r['id']=='NP-480');self.assertEqual(r['values'][12],'');self.assertIn('maps/search/',r['values'][13])
 def test_url_is_encoded(self):
  rows,_=review.prepare(self.fixture());self.assertNotIn('合成',rows[0]['values'][13]);self.assertIn('%',rows[0]['values'][13])
 def test_priority_order(self):self.assertEqual([r['id'] for r in review.prepare(self.fixture())[0][:8]],review.PRIORITY)
 def test_source_values_not_formulas(self):
  d=self.fixture();d['tables']['公園資料'][0]['地址']='=1+1';rows,g=review.prepare(d)
  with tempfile.TemporaryDirectory() as tmp:
   f=Path(tmp)/'review.xlsx';review.make_workbook(d,rows,g,f);w=load_workbook(f);c=next(row[3] for row in w['審核表'].iter_rows(min_row=2) if row[0].value=='NP-1');self.assertEqual(c.data_type,'s');w.close()
 def test_xlsx_round_trip_loader(self):
  d=self.fixture();from openpyxl import Workbook
  with tempfile.TemporaryDirectory() as tmp:
   f=Path(tmp)/'source.xlsx';w=Workbook();w.remove(w.active)
   for name,rows in d['tables'].items():
    sheet=w.create_sheet(name);keys=[k for k in rows[0] if k!='__sheetRow'];sheet.append(keys)
    for r in rows:sheet.append([r[k] for k in keys])
   w.save(f);before=review.sha(f);loaded=review.load_source(f);self.assertEqual(len(review.prepare(loaded)[0]),635);self.assertEqual(review.sha(f),before)
 def test_protected_inputs_locked(self):
  d=self.fixture();rows,g=review.prepare(d)
  with tempfile.TemporaryDirectory() as tmp:
   f=Path(tmp)/'review.xlsx';review.make_workbook(d,rows,g,f);w=load_workbook(f)
   row=next(r for r in w['審核表'].iter_rows(min_row=2) if r[0].value=='NP-9');self.assertTrue(row[16].protection.locked);self.assertTrue(row[17].protection.locked);self.assertEqual(row[23].value,'保護');w.close()
 def test_risk_coordinate_links_clickable(self):
  d=self.fixture();d['tables']['公園座標候選'][0]['候選緯度']=25.02;d['tables']['公園座標候選'][1]['候選緯度']=25.02;rows,g=review.prepare(d)
  with tempfile.TemporaryDirectory() as tmp:
   f=Path(tmp)/'review.xlsx';review.make_workbook(d,rows,g,f);w=load_workbook(f);row=next(r for r in w['高風險優先審核'].iter_rows(min_row=2) if r[0].value=='NP-1');self.assertIsNotNone(row[12].hyperlink);invalid=next(r for r in w['高風險優先審核'].iter_rows(min_row=2) if r[0].value=='NP-480');self.assertIsNone(invalid[12].hyperlink);w.close()
 def test_ods_acceptance_preserves_human_content_and_settings(self):
  from accept_park_review_ods import relabel_copy, STATUS
  import zipfile
  d=self.fixture();rows,g=review.prepare(d)
  with tempfile.TemporaryDirectory() as tmp:
   original=Path(tmp)/'original.xlsx';target=Path(tmp)/'accepted.xlsx'
   w=review.make_workbook(d,rows,g,original);w.close()
   w=load_workbook(original);w['審核表']['W2']='人工備註不得遺失';w['審核表']['O2']='審核中';w.save(original);w.close()
   before=review.sha(original);report=relabel_copy(original,target)
   self.assertTrue(report['allReviewAndHumanInputXmlUnchanged']);self.assertEqual(review.sha(original),before)
   w=load_workbook(target);self.assertEqual(w['審核表']['W2'].value,'人工備註不得遺失');self.assertEqual(w['審核表']['O2'].value,'審核中');self.assertEqual(w['使用說明']['B3'].value,STATUS);self.assertEqual(w.properties.subject,STATUS);w.close()
   with zipfile.ZipFile(original) as a,zipfile.ZipFile(target) as b:
    for n in a.namelist():
     if n not in report['changedMembers']:self.assertEqual(a.read(n),b.read(n))
 def acceptance_fixture(self, tmp):
  d=self.fixture();rows,g=review.prepare(d);source=Path(tmp)/'source.ods';source.write_bytes(b'synthetic mocked ODS');book=Path(tmp)/'review.xlsx';w=review.make_workbook(d,rows,g,book);w.close();return d,source,book
 def test_ods_acceptance_detects_coordinate_difference(self):
  from accept_park_review_ods import reconcile
  from unittest.mock import patch
  with tempfile.TemporaryDirectory() as tmp:
   d,source,book=self.acceptance_fixture(tmp);w=load_workbook(book);w['審核表']['G2']=24.8;w.save(book);w.close()
   with patch('accept_park_review_ods.load_source',return_value=d):_,r=reconcile(source,book)
   self.assertFalse(r['sourcePrimaryOdsVerified']);self.assertTrue(any(x['field']=='候選緯度' for x in r['differences']))
 def test_ods_acceptance_detects_quality_history_difference(self):
  from accept_park_review_ods import reconcile
  from unittest.mock import patch
  with tempfile.TemporaryDirectory() as tmp:
   d,source,book=self.acceptance_fixture(tmp);w=load_workbook(book);w['既有人工修正紀錄']['D2']='altered';w.save(book);w.close()
   with patch('accept_park_review_ods.load_source',return_value=d):_,r=reconcile(source,book)
   self.assertFalse(r['sourcePrimaryOdsVerified']);self.assertTrue(any(x['scope']=='history' for x in r['differences']))
 def test_ods_acceptance_human_edits_not_source_differences(self):
  from accept_park_review_ods import reconcile
  from unittest.mock import patch
  with tempfile.TemporaryDirectory() as tmp:
   d,source,book=self.acceptance_fixture(tmp);w=load_workbook(book);w['審核表']['W2']='人工記錄';w.save(book);w.close();before=review.sha(book)
   with patch('accept_park_review_ods.load_source',return_value=d):_,r=reconcile(source,book)
   self.assertTrue(r['sourcePrimaryOdsVerified']);self.assertEqual(r['differenceCount'],0);self.assertEqual(review.sha(book),before)
 def test_ods_acceptance_mismatch_emits_csv_without_new_workbook(self):
  from accept_park_review_ods import accept
  from unittest.mock import patch
  with tempfile.TemporaryDirectory() as tmp:
   d,source,book=self.acceptance_fixture(tmp);w=load_workbook(book);w['審核表']['D2']='不同地址';w.save(book);w.close();out=Path(tmp)/'result'
   with patch('accept_park_review_ods.load_source',return_value=d):r=accept(source,book,out)
   self.assertFalse(r['sourcePrimaryOdsVerified']);self.assertTrue((out/'source-differences.csv').exists());self.assertFalse((out/book.name).exists())
if __name__=='__main__':unittest.main(verbosity=2)
